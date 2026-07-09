"""Örnek talep ve tedarikçi Excel şablonlarını üretir.

Talep örneği, planlamadan gelen e-postadaki tablo düzenindedir.
Tedarikçi listesi hayali verilerle doldurulmuştur; kendi onaylı
tedarikçilerinizle değiştirin.
"""

from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

_INCE = Side(style="thin")
_CERCEVE = Border(left=_INCE, right=_INCE, top=_INCE, bottom=_INCE)
_BASLIK_DOLGU = PatternFill("solid", fgColor="D9E1F2")

TALEP_ORNEGI = [
    ("HMPT000220", "DOTP", 20000, "KG.", "17.06.2026"),
    ("HMDH000015", "MEG", 2000, "KG.", "17.06.2026"),
    ("HMDH000073", "AKRİLİK KALINLAŞTIRICI", 1000, "KG.", "19.06.2026"),
    ("HMDH000018", "AKRİLİK BİYOSİT", 400, "KG.", "19.06.2026"),
    ("HMDH000100", "ISLATICI PRO", 1000, "KG.", "18.06.2026"),
    ("HMDH000016", "AMONYAK", 300, "KG.", "18.06.2026"),
    ("HMHR000223", "PİGMENT PASTA OKSİT KIRMIZI DOLPHIN", 15, "KG.", "16.06.2026"),
    ("HMHR000248", "PİGMENT PASTA OKSİT KAHVERENGİ DOLPHIN", 25, "KG.", "16.06.2026"),
    ("PMES200008", "0526 BOYA KAPAK ETİKETİ MAT SİYAH R9005", 60000, "ADET", "16.06.2026"),
    ("PMLB002312", "BASKISIZ BEYAZ KARTUŞ KOLİSİ 12'Lİ", 2000, "ADET", "16.06.2026"),
]

TEDARIKCI_ORNEGI = [
    ("HMPT000220", "DOTP", "ABC Kimya A.Ş.", "satis@abckimya.example.com", "Ahmet Yılmaz"),
    ("HMPT000220", "DOTP", "Delta Plastifiyan Ltd.", "teklif@deltaplast.example.com", ""),
    ("HMDH000015", "MEG", "ABC Kimya A.Ş.", "satis@abckimya.example.com", "Ahmet Yılmaz"),
    ("HMDH000015", "MEG", "Ege Solvent San. Tic.", "info@egesolvent.example.com", "Zeynep Kaya"),
    ("HMDH000073", "AKRİLİK KALINLAŞTIRICI", "Polimer Teknik A.Ş.", "satis@polimerteknik.example.com", ""),
    ("HMDH000018", "AKRİLİK BİYOSİT", "Polimer Teknik A.Ş.", "satis@polimerteknik.example.com", ""),
    ("HMDH000018", "AKRİLİK BİYOSİT", "BioKim Ltd.", "siparis@biokim.example.com", "Murat Demir"),
    ("HMDH000100", "ISLATICI PRO", "Polimer Teknik A.Ş.", "satis@polimerteknik.example.com", ""),
    ("HMDH000016", "AMONYAK", "Ege Solvent San. Tic.", "info@egesolvent.example.com", "Zeynep Kaya"),
    ("HMHR000223", "PİGMENT PASTA OKSİT KIRMIZI DOLPHIN", "Dolphin Pigment A.Ş.", "satis@dolphinpigment.example.com", ""),
    ("HMHR000248", "PİGMENT PASTA OKSİT KAHVERENGİ DOLPHIN", "Dolphin Pigment A.Ş.", "satis@dolphinpigment.example.com", ""),
    ("PMES200008", "0526 BOYA KAPAK ETİKETİ MAT SİYAH R9005", "Yıldız Etiket Matbaa", "teklif@yildizetiket.example.com", "Elif Şahin"),
    ("PMES200008", "0526 BOYA KAPAK ETİKETİ MAT SİYAH R9005", "Marmara Baskı Ltd.", "satis@marmarabaski.example.com", ""),
    ("PMLB002312", "BASKISIZ BEYAZ KARTUŞ KOLİSİ 12'Lİ", "Anadolu Ambalaj A.Ş.", "siparis@anadoluambalaj.example.com", ""),
    ("PMLB002312", "BASKISIZ BEYAZ KARTUŞ KOLİSİ 12'Lİ", "Marmara Baskı Ltd.", "satis@marmarabaski.example.com", ""),
]


def _tablo_yaz(ws, basliklar, satirlar, genislikler):
    for i, b in enumerate(basliklar, start=1):
        h = ws.cell(1, i, b)
        h.font = Font(bold=True)
        h.fill = _BASLIK_DOLGU
        h.border = _CERCEVE
        h.alignment = Alignment(horizontal="center")
    for r, satir in enumerate(satirlar, start=2):
        for i, deger in enumerate(satir, start=1):
            ws.cell(r, i, deger).border = _CERCEVE
    for i, g in enumerate(genislikler, start=1):
        ws.column_dimensions[chr(64 + i)].width = g


def ornekleri_olustur(klasor: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Talep"
    _tablo_yaz(ws, ["Kodu", "İsmi", "Miktar", "Birim", "Teslim tarihi"], TALEP_ORNEGI, [14, 42, 12, 8, 14])
    for r in range(2, len(TALEP_ORNEGI) + 2):
        ws.cell(r, 3).number_format = "#,##0.00"
    wb.save(os.path.join(klasor, "ornek_talep.xlsx"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Tedarikçiler"
    _tablo_yaz(
        ws,
        ["Stok Kodu", "Stok Adı", "Tedarikçi Adı", "E-posta", "Yetkili"],
        TEDARIKCI_ORNEGI,
        [14, 42, 28, 36, 18],
    )
    wb.save(os.path.join(klasor, "ornek_tedarikciler.xlsx"))


if __name__ == "__main__":
    ornekleri_olustur(".")
