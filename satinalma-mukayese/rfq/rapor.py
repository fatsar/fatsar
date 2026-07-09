"""Toplanan tekliflerden mukayese (karşılaştırma) raporu üretimi."""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .form import Teklif, TeklifKalemi
from .veri import TalepKalemi

_INCE = Side(style="thin")
_CERCEVE = Border(left=_INCE, right=_INCE, top=_INCE, bottom=_INCE)
_BASLIK_DOLGU = PatternFill("solid", fgColor="D9E1F2")
_EN_UCUZ_DOLGU = PatternFill("solid", fgColor="C6EFCE")
_EN_UCUZ_YAZI = Font(bold=True, color="006100")
_YOK_DOLGU = PatternFill("solid", fgColor="F2F2F2")


def mukayese_raporu(dosya: str, rfq_no: str, kalemler: list[TalepKalemi], teklifler: list[Teklif]) -> None:
    """Kalem satırları x tedarikçi sütunları düzeninde mukayese tablosu yazar.

    Her tedarikçi için Birim Fiyat / Tutar / Termin / Vade sütunları açılır;
    kalem bazında en düşük birim fiyat yeşil vurgulanır.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Mukayese"

    tedarikciler = sorted({t.tedarikci or t.kaynak for t in teklifler})
    fiyatlar: dict[tuple[str, str], TeklifKalemi] = {}
    for teklif in teklifler:
        ad = teklif.tedarikci or teklif.kaynak
        for tk in teklif.kalemler:
            fiyatlar[(ad, tk.kod)] = tk

    ws["A1"] = f"TEKLİF MUKAYESE TABLOSU — {rfq_no}"
    ws["A1"].font = Font(bold=True, size=14)

    sabit_basliklar = ["Stok Kodu", "Stok Adı", "Miktar", "Birim", "İstenen Teslim"]
    ted_basliklar = ["Birim Fiyat", "Para B.", "Tutar", "Termin", "Vade"]

    ust, alt = 3, 4
    for i, b in enumerate(sabit_basliklar, start=1):
        ws.merge_cells(start_row=ust, start_column=i, end_row=alt, end_column=i)
        h = ws.cell(ust, i, b)
        h.font = Font(bold=True)
        h.fill = _BASLIK_DOLGU
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sutun = len(sabit_basliklar) + 1
    ted_sutunlari: dict[str, int] = {}
    for ad in tedarikciler:
        ted_sutunlari[ad] = sutun
        ws.merge_cells(start_row=ust, start_column=sutun, end_row=ust, end_column=sutun + len(ted_basliklar) - 1)
        h = ws.cell(ust, sutun, ad)
        h.font = Font(bold=True)
        h.fill = _BASLIK_DOLGU
        h.alignment = Alignment(horizontal="center")
        for j, b in enumerate(ted_basliklar):
            h2 = ws.cell(alt, sutun + j, b)
            h2.font = Font(bold=True)
            h2.fill = _BASLIK_DOLGU
            h2.alignment = Alignment(horizontal="center")
        sutun += len(ted_basliklar)
    son_sutun = sutun - 1
    for r in (ust, alt):
        for c in range(1, son_sutun + 1):
            ws.cell(r, c).border = _CERCEVE

    toplamlar: dict[str, float] = {ad: 0.0 for ad in tedarikciler}
    eksiksiz: dict[str, bool] = {ad: True for ad in tedarikciler}

    satir = alt + 1
    for kalem in kalemler:
        ws.cell(satir, 1, kalem.kod)
        ws.cell(satir, 2, kalem.isim)
        ws.cell(satir, 3, kalem.miktar).number_format = "#,##0.00"
        ws.cell(satir, 4, kalem.birim)
        ws.cell(satir, 5, kalem.teslim_tarihi)

        verilenler = {
            ad: fiyatlar[(ad, kalem.kod)] for ad in tedarikciler if (ad, kalem.kod) in fiyatlar
        }
        en_dusuk = min((tk.birim_fiyat for tk in verilenler.values()), default=None)

        for ad in tedarikciler:
            c = ted_sutunlari[ad]
            tk = verilenler.get(ad)
            if tk is None:
                for j in range(len(ted_basliklar)):
                    ws.cell(satir, c + j, "-").fill = _YOK_DOLGU
                eksiksiz[ad] = False
            else:
                f_hucre = ws.cell(satir, c, tk.birim_fiyat)
                f_hucre.number_format = "#,##0.00"
                ws.cell(satir, c + 1, tk.para_birimi)
                t_hucre = ws.cell(satir, c + 2, tk.miktar * tk.birim_fiyat)
                t_hucre.number_format = "#,##0.00"
                ws.cell(satir, c + 3, tk.termin)
                ws.cell(satir, c + 4, tk.odeme_vadesi)
                toplamlar[ad] += tk.miktar * tk.birim_fiyat
                if en_dusuk is not None and tk.birim_fiyat == en_dusuk:
                    f_hucre.fill = _EN_UCUZ_DOLGU
                    f_hucre.font = _EN_UCUZ_YAZI
        for c in range(1, son_sutun + 1):
            ws.cell(satir, c).border = _CERCEVE
        satir += 1

    # Toplam satırı
    ws.cell(satir, 1, "TOPLAM").font = Font(bold=True)
    dolu_toplamlar = [toplamlar[ad] for ad in tedarikciler if toplamlar[ad] > 0]
    en_dusuk_toplam = min(dolu_toplamlar, default=None)
    for ad in tedarikciler:
        c = ted_sutunlari[ad]
        h = ws.cell(satir, c + 2, toplamlar[ad] if toplamlar[ad] > 0 else "-")
        h.number_format = "#,##0.00"
        h.font = Font(bold=True)
        if en_dusuk_toplam is not None and toplamlar[ad] == en_dusuk_toplam:
            h.fill = _EN_UCUZ_DOLGU
            h.font = _EN_UCUZ_YAZI
        if not eksiksiz[ad]:
            ws.cell(satir + 1, c, "(eksik kalem var)").font = Font(italic=True, size=9)
    for c in range(1, son_sutun + 1):
        ws.cell(satir, c).border = _CERCEVE

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 11
    ws.column_dimensions["D"].width = 7
    ws.column_dimensions["E"].width = 14
    from openpyxl.utils import get_column_letter

    for c in range(len(sabit_basliklar) + 1, son_sutun + 1):
        ws.column_dimensions[get_column_letter(c)].width = 12

    ws.freeze_panes = "F5"
    wb.save(dosya)
