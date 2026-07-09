"""Tedarikçiye gönderilen teklif formunun oluşturulması ve okunması."""

from __future__ import annotations

from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .veri import TalepKalemi, Tedarikci, sayi_cevir, tarih_metni

FORM_BASLIKLARI = [
    "Stok Kodu",
    "Stok Adı",
    "Miktar",
    "Birim",
    "İstenen Teslim Tarihi",
    "Birim Fiyat",
    "Para Birimi",
    "Teslim Tarihiniz (Termin)",
    "Ödeme Vadesi",
    "Not",
]

_INCE = Side(style="thin")
_CERCEVE = Border(left=_INCE, right=_INCE, top=_INCE, bottom=_INCE)
_BASLIK_DOLGU = PatternFill("solid", fgColor="D9E1F2")
_DOLDUR_DOLGU = PatternFill("solid", fgColor="FFF2CC")


@dataclass
class TeklifKalemi:
    kod: str
    isim: str
    miktar: float
    birim: str
    birim_fiyat: float
    para_birimi: str
    termin: str = ""
    odeme_vadesi: str = ""
    not_: str = ""

    @property
    def tutar(self) -> float:
        return self.miktar * self.birim_fiyat


@dataclass
class Teklif:
    rfq_no: str
    tedarikci: str
    kalemler: list[TeklifKalemi] = field(default_factory=list)
    kaynak: str = ""  # dosya yolu veya e-posta bilgisi


def teklif_formu_olustur(dosya: str, rfq_no: str, tedarikci: Tedarikci, kalemler: list[TalepKalemi]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Teklif Formu"

    ws["A1"] = "TEKLİF FORMU"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Teklif No:"
    ws["B2"] = rfq_no
    ws["A3"] = "Tedarikçi:"
    ws["B3"] = tedarikci.ad
    ws["A4"] = (
        "Lütfen sarı alanları doldurup bu dosyayı e-postaya EK olarak, "
        "konu satırını değiştirmeden yanıtlayınız."
    )
    ws["A4"].font = Font(italic=True, color="C00000")
    for satir in (2, 3):
        ws.cell(satir, 1).font = Font(bold=True)

    baslik_satiri = 6
    for i, baslik in enumerate(FORM_BASLIKLARI, start=1):
        hucre = ws.cell(baslik_satiri, i, baslik)
        hucre.font = Font(bold=True)
        hucre.fill = _BASLIK_DOLGU
        hucre.border = _CERCEVE
        hucre.alignment = Alignment(horizontal="center", wrap_text=True)

    for j, kalem in enumerate(kalemler):
        satir = baslik_satiri + 1 + j
        degerler = [kalem.kod, kalem.isim, kalem.miktar, kalem.birim, kalem.teslim_tarihi, None, "TL", None, None, None]
        for i, deger in enumerate(degerler, start=1):
            hucre = ws.cell(satir, i, deger)
            hucre.border = _CERCEVE
            if i >= 6:  # tedarikçinin dolduracağı alanlar
                hucre.fill = _DOLDUR_DOLGU
        ws.cell(satir, 3).number_format = "#,##0.00"
        ws.cell(satir, 6).number_format = "#,##0.00"

    genislikler = [14, 42, 12, 8, 18, 12, 12, 18, 14, 24]
    for i, g in enumerate(genislikler, start=1):
        ws.column_dimensions[chr(64 + i)].width = g

    wb.save(dosya)


def teklif_formu_oku(dosya: str) -> Teklif:
    """Tedarikçiden dönen doldurulmuş formu okur. Fiyatsız satırlar atlanır."""
    wb = load_workbook(dosya, data_only=True)
    ws = wb.active

    rfq_no, tedarikci = "", ""
    baslik_satiri = None
    for satir in range(1, min(ws.max_row, 15) + 1):
        ilk = str(ws.cell(satir, 1).value or "").strip()
        if ilk.lower().startswith("teklif no"):
            rfq_no = str(ws.cell(satir, 2).value or "").strip()
        elif ilk.lower().startswith("tedarikçi") or ilk.lower().startswith("tedarikci"):
            tedarikci = str(ws.cell(satir, 2).value or "").strip()
        elif ilk == FORM_BASLIKLARI[0]:
            baslik_satiri = satir
            break
    if baslik_satiri is None:
        raise ValueError(f"'{dosya}' teklif formu düzeninde değil (başlık satırı yok).")

    teklif = Teklif(rfq_no=rfq_no, tedarikci=tedarikci, kaynak=dosya)
    for satir in range(baslik_satiri + 1, ws.max_row + 1):
        kod = ws.cell(satir, 1).value
        if kod is None or not str(kod).strip():
            continue
        fiyat_ham = ws.cell(satir, 6).value
        if fiyat_ham is None or str(fiyat_ham).strip() == "":
            continue  # fiyat verilmemiş kalem
        try:
            fiyat = sayi_cevir(fiyat_ham)
            miktar = sayi_cevir(ws.cell(satir, 3).value or 0)
        except ValueError:
            continue
        teklif.kalemler.append(
            TeklifKalemi(
                kod=str(kod).strip(),
                isim=str(ws.cell(satir, 2).value or "").strip(),
                miktar=miktar,
                birim=str(ws.cell(satir, 4).value or "").strip(),
                birim_fiyat=fiyat,
                para_birimi=str(ws.cell(satir, 7).value or "TL").strip() or "TL",
                termin=tarih_metni(ws.cell(satir, 8).value),
                odeme_vadesi=str(ws.cell(satir, 9).value or "").strip(),
                not_=str(ws.cell(satir, 10).value or "").strip(),
            )
        )
    return teklif
