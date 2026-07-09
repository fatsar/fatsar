"""Excel dosyalarından talep ve tedarikçi verilerini okuma."""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field

from openpyxl import load_workbook

# Başlık satırını bulmak için kabul edilen sütun adları (küçük harf, noktalama yok)
TALEP_KOLONLARI = {
    "kod": {"kodu", "kod", "stok kodu", "stokkodu", "malzeme kodu"},
    "isim": {"ismi", "isim", "stok adi", "stok adı", "malzeme adi", "malzeme adı", "aciklama", "açıklama"},
    "miktar": {"miktar", "miktarı", "adet"},
    "birim": {"birim", "olcu birimi", "ölçü birimi"},
    "teslim": {"teslim tarihi", "teslim", "termin", "termin tarihi", "istenen tarih"},
}

TEDARIKCI_KOLONLARI = {
    "kod": TALEP_KOLONLARI["kod"],
    "isim": TALEP_KOLONLARI["isim"],
    "tedarikci": {"tedarikci", "tedarikçi", "tedarikci adi", "tedarikçi adı", "firma", "firma adi", "firma adı"},
    "eposta": {"e-posta", "eposta", "e posta", "email", "e-mail", "mail"},
    "yetkili": {"yetkili", "ilgili", "kontak", "yetkili kisi", "yetkili kişi"},
}


@dataclass
class TalepKalemi:
    kod: str
    isim: str
    miktar: float
    birim: str
    teslim_tarihi: str

    def sozluk(self) -> dict:
        return {
            "kod": self.kod,
            "isim": self.isim,
            "miktar": self.miktar,
            "birim": self.birim,
            "teslim_tarihi": self.teslim_tarihi,
        }


@dataclass
class Tedarikci:
    ad: str
    eposta: str
    yetkili: str = ""
    kalemler: list = field(default_factory=list)  # bu tedarikçiden istenecek TalepKalemi listesi


def _normalize(deger) -> str:
    if deger is None:
        return ""
    return str(deger).strip().lower().replace("i̇", "i")


def _basligi_bul(ws, kolonlar: dict) -> tuple[int, dict]:
    """İlk 10 satırda başlık satırını arar; (satır no, {alan: sütun no}) döner."""
    for satir in range(1, min(ws.max_row, 10) + 1):
        eslesme = {}
        for sutun in range(1, ws.max_column + 1):
            hucre = _normalize(ws.cell(satir, sutun).value)
            if not hucre:
                continue
            for alan, adaylar in kolonlar.items():
                if alan not in eslesme and hucre in adaylar:
                    eslesme[alan] = sutun
        # kod + bir alan daha bulunduysa başlık kabul et
        if "kod" in eslesme and len(eslesme) >= 2:
            return satir, eslesme
    raise ValueError(
        "Başlık satırı bulunamadı. Excel'de 'Kodu', 'İsmi', 'Miktar' gibi "
        "başlıkların olduğu bir satır olmalı."
    )


def sayi_cevir(deger) -> float:
    """'20.000,00' gibi Türkçe biçimli sayıları da çevirir."""
    if deger is None or deger == "":
        raise ValueError("Boş sayı")
    if isinstance(deger, (int, float)):
        return float(deger)
    metin = str(deger).strip().replace(" ", "")
    if "," in metin:
        metin = metin.replace(".", "").replace(",", ".")
    return float(metin)


def tarih_metni(deger) -> str:
    if deger is None:
        return ""
    if isinstance(deger, (dt.datetime, dt.date)):
        return deger.strftime("%d.%m.%Y")
    return str(deger).strip()


def talep_oku(dosya: str) -> list[TalepKalemi]:
    wb = load_workbook(dosya, data_only=True)
    ws = wb.active
    baslik, sutunlar = _basligi_bul(ws, TALEP_KOLONLARI)
    kalemler = []
    for satir in range(baslik + 1, ws.max_row + 1):
        kod = ws.cell(satir, sutunlar["kod"]).value
        if kod is None or not str(kod).strip():
            continue
        try:
            miktar = sayi_cevir(ws.cell(satir, sutunlar["miktar"]).value) if "miktar" in sutunlar else 0.0
        except ValueError:
            continue
        kalemler.append(
            TalepKalemi(
                kod=str(kod).strip(),
                isim=str(ws.cell(satir, sutunlar["isim"]).value or "").strip() if "isim" in sutunlar else "",
                miktar=miktar,
                birim=str(ws.cell(satir, sutunlar["birim"]).value or "").strip() if "birim" in sutunlar else "",
                teslim_tarihi=tarih_metni(ws.cell(satir, sutunlar["teslim"]).value) if "teslim" in sutunlar else "",
            )
        )
    if not kalemler:
        raise ValueError(f"'{dosya}' içinde talep kalemi bulunamadı.")
    return kalemler


def tedarikci_oku(dosya: str) -> dict[str, list[Tedarikci]]:
    """Stok kodu -> onaylı tedarikçi listesi eşlemesi döner."""
    wb = load_workbook(dosya, data_only=True)
    ws = wb.active
    baslik, sutunlar = _basligi_bul(ws, TEDARIKCI_KOLONLARI)
    if "tedarikci" not in sutunlar or "eposta" not in sutunlar:
        raise ValueError(
            "Tedarikçi listesinde 'Tedarikçi' ve 'E-posta' sütunları bulunamadı."
        )
    eslem: dict[str, list[Tedarikci]] = {}
    for satir in range(baslik + 1, ws.max_row + 1):
        kod = ws.cell(satir, sutunlar["kod"]).value
        ad = ws.cell(satir, sutunlar["tedarikci"]).value
        eposta = ws.cell(satir, sutunlar["eposta"]).value
        if not kod or not ad or not eposta:
            continue
        yetkili = ""
        if "yetkili" in sutunlar:
            yetkili = str(ws.cell(satir, sutunlar["yetkili"]).value or "").strip()
        eslem.setdefault(str(kod).strip(), []).append(
            Tedarikci(ad=str(ad).strip(), eposta=str(eposta).strip(), yetkili=yetkili)
        )
    if not eslem:
        raise ValueError(f"'{dosya}' içinde tedarikçi kaydı bulunamadı.")
    return eslem


def tedarikcilere_dagit(
    kalemler: list[TalepKalemi], eslem: dict[str, list[Tedarikci]]
) -> tuple[dict[str, Tedarikci], list[TalepKalemi]]:
    """Talep kalemlerini onaylı tedarikçilere dağıtır.

    Aynı tedarikçi birden çok kalemde onaylıysa tek e-postada birleştirilir.
    (tedarikçiler sözlüğü, tedarikçisi bulunamayan kalemler) döner.
    """
    tedarikciler: dict[str, Tedarikci] = {}
    eksikler: list[TalepKalemi] = []
    for kalem in kalemler:
        adaylar = eslem.get(kalem.kod, [])
        if not adaylar:
            eksikler.append(kalem)
            continue
        for aday in adaylar:
            anahtar = aday.eposta.lower()
            if anahtar not in tedarikciler:
                tedarikciler[anahtar] = Tedarikci(ad=aday.ad, eposta=aday.eposta, yetkili=aday.yetkili)
            tedarikciler[anahtar].kalemler.append(kalem)
    return tedarikciler, eksikler
